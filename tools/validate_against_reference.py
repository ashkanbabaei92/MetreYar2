# -*- coding: utf-8 -*-
# validate_against_reference.py — اعتبارسنجیِ خروجی‌ها در برابر اعدادِ مرجع.
#
# دو حالت دارد:
#   python tools/validate_against_reference.py              # اجرای موتور + اعتبارسنجی
#   python tools/validate_against_reference.py --xlsx FILE  # اعتبارسنجیِ فایل Excel
#
# تلرانسِ پیش‌فرض: ۱ ریال (قابل تغییر با --tolerance).
#
# اعدادِ مرجع از خطِ لولهٔ کاملِ پروژهٔ مرجع استخراج شده‌اند و با آزمون‌های
# T112/T113 در هر دو زبان قفل شده‌اند. این ابزار یک بررسیِ مستقلِ نهایی
# است — اگر موتور یا Excel تغییر کند و این اعداد جابجا شوند، خطا می‌دهد.
"""اعتبارسنجیِ اعدادِ مرجعِ متره‌یار (موتور یا فایلِ Excel)."""
import argparse
import os
import re
import sys

# خروجیِ کنسولِ ویندوز cp1252 است؛ UTF-8 می‌کنیم تا پیام‌های فارسی چاپ شوند.
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

# ارقامِ فارسی برای تبدیلِ رشته‌های قالب‌بندی‌شده (داخلِ app.py هم هست)
FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")

ROOT = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)

# ── اعدادِ مرجعِ پروژهٔ نمونه (تنظیماتِ پیش‌فرض، It = ۱۴۰۴Q4) ──
# منبع: tests/char_python/test_T112_quality_parity.py و test_T113...
EXPECTED = {
    "base_total": 23_218_381_337.01,
    "grand_total": 31_391_251_567.64,
    "baravord_q4": 45_437_742_820.32,
}

# تلرانسِ پیش‌فرض: ۱ ریال
DEFAULT_TOLERANCE = 1.0


def check_number(name, actual, expected, tolerance=DEFAULT_TOLERANCE):
    """مقایسهٔ یک عدد با مرجع. None اگر درست، در غیر این صورت پیامِ خطا.

    None (مقدارِ ناموجود) همیشه خطا است — عددِ مرجع باید همیشه موجود باشد.
    """
    if actual is None:
        return "%s: موجود نیست (انتظار: %.2f)" % (name, expected)
    diff = abs(float(actual) - expected)
    if diff > tolerance:
        return ("%s: %.2f ≠ %.2f (اختلاف: %.2f ریال، تلرانس: %.2f)"
                % (name, float(actual), expected, diff, tolerance))
    return None


def validate_numbers(actuals, expected=None, tolerance=DEFAULT_TOLERANCE):
    """اعتبارسنجیِ یک دیکشنریِ اعداد در برابر مرجع. لیستِ خطاها (خالی = OK)."""
    expected = expected if expected is not None else EXPECTED
    issues = []
    for name, exp in expected.items():
        msg = check_number(name, actuals.get(name), exp, tolerance)
        if msg:
            issues.append(msg)
    return issues


def run_engine():
    """اجرای خطِ لولهٔ کاملِ پروژهٔ مرجع → دیکشنریِ اعداد."""
    from app import (DynamoMappingParser, RevitScheduleParser,
                     TechnicalOfficeEngine, _default_fehrest_db,
                     default_settings)

    revit_dir = os.path.join(ROOT, "data", "export revit")
    dynamo_dir = os.path.join(ROOT, "data", "dynamo")

    tables = []
    for fname in sorted(os.listdir(revit_dir)):
        if fname.endswith(".txt"):
            with open(os.path.join(revit_dir, fname), "r",
                      encoding="utf-8-sig") as fh:
                tables.append(RevitScheduleParser.parse(fname, fh.read()))
    results = []
    for root, _dirs, fnames in os.walk(dynamo_dir):
        for fn in sorted(fnames):
            if fn.endswith(".txt"):
                with open(os.path.join(root, fn), "rb") as fh:
                    results.append(
                        DynamoMappingParser.auto_detect_and_parse(fh.read(), fn))
    map_df = DynamoMappingParser.merge_formats(results)

    s = default_settings()
    s.update({"baravord_enabled": True, "baravord_approved": True})
    result = TechnicalOfficeEngine(_default_fehrest_db(), s).compute(tables,
                                                                    map_df)
    bar = (result.get("adjustment") or {}).get("baravord") or {}
    return {
        "base_total": result.get("base_total"),
        "grand_total": result.get("grand_total"),
        "baravord_q4": bar.get("result_amount"),
    }


def _cell_numbers(v):
    """اعدادِ موجود در یک سلول را برمی‌گرداند (عددیِ خالص یا رشتهٔ فارسی).

    شیتِ «برآورد به‌هنگام» مبالغ را به‌صورتِ رشته‌های قالب‌بندی‌شدهٔ فارسی
    می‌نویسد (مثل «۴۵,۴۳۷,۷۴۲,۸۲۰»)؛ شیت‌های دیگر سلولِ عددی دارند. هر دو
    حالت پشتیبانی می‌شود.
    """
    if isinstance(v, bool):
        return []
    if isinstance(v, (int, float)):
        return [float(v)]
    if isinstance(v, str):
        out = []
        # توکن‌های عددی: ارقامِ فارسی یا انگلیسی با ویرگول/نقطه
        for tok in re.findall(r"[۰-۹0-9][۰-۹0-9,.\-]*", v):
            en = tok.translate(FA_DIGITS)
            cleaned = en.replace(",", "")
            try:
                out.append(float(cleaned))
            except ValueError:
                pass
        return out
    return []


def validate_xlsx(path, expected=None, tolerance=DEFAULT_TOLERANCE):
    """اعتبارسنجیِ فایلِ Excel: هر عددِ مرجع باید حداقل در یک سلول دیده شود.

    سلول‌های عددی و رشته‌های قالب‌بندی‌شدهٔ فارسی (شیتِ برآورد به‌هنگام)
    هر دو بررسی می‌شوند. کفهٔ راست: اگر عددی در کلِ کاربرگ نیابد، یعنی یا
    موتور اشتباه است یا برچسب/فرمتِ ستون جابجا شده.
    """
    from openpyxl import load_workbook

    expected = expected if expected is not None else EXPECTED
    wb = load_workbook(path, data_only=True)
    found = {name: False for name in expected}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                for num in _cell_numbers(cell.value):
                    for name, exp in expected.items():
                        if abs(num - exp) <= tolerance:
                            found[name] = True
    wb.close()
    issues = []
    for name, ok in found.items():
        if not ok:
            issues.append("%s: در هیچ سلولی از %s یافت نشد (انتظار: %.2f)"
                          % (name, os.path.basename(path), expected[name]))
    return issues


def main(argv=None):
    p = argparse.ArgumentParser(
        description="اعتبارسنجیِ اعدادِ مرجعِ متره‌یار")
    p.add_argument("--xlsx", metavar="FILE",
                   help="به‌جای اجرای موتور، این فایل Excel را بررسی کن")
    p.add_argument("--tolerance", type=float, default=DEFAULT_TOLERANCE,
                   help="تلرانس بر حسب ریال (پیش‌فرض: %.2f)" % DEFAULT_TOLERANCE)
    args = p.parse_args(argv)

    if args.xlsx:
        if not os.path.isfile(args.xlsx):
            print("خطا: فایل پیدا نشد: %s" % args.xlsx)
            return 2
        print("بررسیِ Excel: %s (تلرانس: %.2f ریال)" % (args.xlsx, args.tolerance))
        issues = validate_xlsx(args.xlsx, tolerance=args.tolerance)
    else:
        print("اجرای موتور روی پروژهٔ مرجع ...")
        actuals = run_engine()
        print("  base_total  = %.2f" % actuals["base_total"])
        print("  grand_total = %.2f" % actuals["grand_total"])
        print("  baravord_q4 = %.2f" % actuals["baravord_q4"])
        print("اعتبارسنجی (تلرانس: %.2f ریال) ..." % args.tolerance)
        issues = validate_numbers(actuals, tolerance=args.tolerance)

    if issues:
        print("")
        for msg in issues:
            print("  FAIL  %s" % msg)
        print("")
        print("%d مورد نامطبوع." % len(issues))
        return 1
    print("")
    print("همهٔ اعدادِ مرجع در تلرانس هستند. OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
