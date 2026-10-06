# -*- coding: utf-8 -*-
# T132 — حالتِ ارائه (فاز E-10.1 §۳).
#
# پرچمِ demo_mode واترمارکِ «نسخهٔ نمایشی – اعتبارسنجی‌نشده با صورت‌وضعیت
# تأییدشده» را روی همهٔ برگه‌های Excel و همهٔ صفحاتِ UI می‌نشاند. این
# آزمون قفل می‌کند:
#   ۱) demo_mode=True → واترمارک در همهٔ شیت‌های Excel هست؛
#   ۲) demo_mode=False → واترمارک در هیچ شیتی نیست؛
#   ۳) واترمارکِ دمو و واترمارکِ «پیش‌نویس» کنار هم کار می‌کنند؛
#   ۴) پروژهٔ نمونه (examples/sample_project) در demo_mode تحویل می‌شود.
# آینهٔ test_T132_demo_mode.ts
import io
import json
import os
import sys

import pytest

pytestmark = pytest.mark.slow

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from app import (DEMO_BANNER, DynamoMappingParser, ExcelExporter,
                 RevitScheduleParser, TechnicalOfficeEngine,
                 _default_fehrest_db, default_settings)

ROOT = os.path.join(os.path.dirname(__file__), "..", "..")
REVIT_DIR = os.path.join(ROOT, "data", "export revit")
DYNAMO_DIR = os.path.join(ROOT, "data", "dynamo")
SAMPLE_DIR = os.path.join(ROOT, "examples", "sample_project")

DRAFT_BANNER = "پیش‌نویس – دورهٔ کارکرد وارد نشده"

META = {"report_type": "دوره‌ای", "report_no": "۱", "contract_no": "۱",
        "project_name": "پروژه نمونه", "employer": "کارفرما", "contractor": "پیمانکار",
        "consultant": "مشاور", "period_start": "۱۴۰۴/۰۱/۰۱", "period_end": "۱۴۰۴/۰۳/۳۱",
        "location": "گرگان", "contract_date": "۱۴۰۴/۰۱/۰۱"}

_XLS_EXTRA = dict(tadil_enabled=True, baravord_enabled=True,
                  proposal_date="۱۴۰۳/۱۲/۲۹",
                  work_period_start="۱۴۰۴/۰۱/۰۱",
                  work_period_end="۱۴۰۵/۰۲/۱۵",
                  provisional_coef_1405q1=1.1530518586507572,
                  completed_within_term=True,
                  contract_status="فسخ با قصور")


def _run_engine(**over):
    tables = []
    for fname in sorted(os.listdir(REVIT_DIR)):
        if fname.endswith(".txt"):
            with open(os.path.join(REVIT_DIR, fname), "r", encoding="utf-8-sig") as fh:
                tables.append(RevitScheduleParser.parse(fname, fh.read()))
    results = []
    for root, _dirs, fnames in os.walk(DYNAMO_DIR):
        for fn in sorted(fnames):
            if fn.endswith(".txt"):
                with open(os.path.join(root, fn), "rb") as fh:
                    results.append(
                        DynamoMappingParser.auto_detect_and_parse(fh.read(), fn))
    map_df = DynamoMappingParser.merge_formats(results)
    s = default_settings()
    s.update(_XLS_EXTRA)
    s.update(over)
    return TechnicalOfficeEngine(_default_fehrest_db(), s).compute(tables, map_df), s


def _sheet_values(buf):
    """مقادیرِ همهٔ سلول‌های هر برگه (با نام برگه)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(buf), data_only=False)
    return {ws.title: [c.value for row in ws.iter_rows() for c in row
                       if c.value is not None]
            for ws in wb.worksheets}


def _sheets_without(values, needle):
    return [name for name, vals in values.items()
            if not any(isinstance(v, str) and needle in v for v in vals)]


def test_T132_banner_text_is_pinned():
    """متنِ واترمارک دقیقاً همین باشد (یک غلطِ املایی فوراً خودش را نشان دهد)."""
    assert DEMO_BANNER == "نسخهٔ نمایشی – اعتبارسنجی‌نشده با صورت‌وضعیت تأییدشده"


def test_T132_demo_watermark_on_all_sheets():
    res, s = _run_engine()
    buf = ExcelExporter().export(META, res, {**s, "demo_mode": True}, [])
    values = _sheet_values(buf)
    assert values, "خروجی باید برگه داشته باشد"
    missing = _sheets_without(values, DEMO_BANNER)
    assert not missing, "شیتی بدون واترمارک: %s" % missing


def test_T132_no_watermark_when_demo_off():
    res, s = _run_engine()
    buf = ExcelExporter().export(META, res, {**s, "demo_mode": False}, [])
    hit = [name for name, vals in _sheet_values(buf).items()
           if any(isinstance(v, str) and DEMO_BANNER in v for v in vals)]
    assert not hit, "واترمارکِ دمو در حالتِ عادی نباید باشد: %s" % hit


def test_T132_demo_and_draft_banners_coexist():
    """دورهٔ کارکردِ خالی + demo_mode: هر دو واترمارک روی همهٔ شیت‌ها."""
    res, s = _run_engine(work_period_start="", work_period_end="")
    meta = {**META, "report_type": "پیش‌نویس"}
    buf = ExcelExporter().export(meta, res, {**s, "demo_mode": True}, [])
    values = _sheet_values(buf)
    assert not _sheets_without(values, DEMO_BANNER), "واترمارکِ دمو"
    assert not _sheets_without(values, DRAFT_BANNER), "واترمارکِ پیش‌نویس"


def test_T132_sample_project_is_demo():
    """پروژهٔ نمونه در حالتِ نمایشی تحویل می‌شود (پروفایل + Excel)."""
    with open(os.path.join(SAMPLE_DIR, "project_profile.json"), "r",
              encoding="utf-8") as fh:
        prof = json.load(fh)
    assert prof["settings"]["demo_mode"] is True
    with open(os.path.join(SAMPLE_DIR, "sample_output.xlsx"), "rb") as fh:
        buf = fh.read()
    missing = _sheets_without(_sheet_values(buf), DEMO_BANNER)
    assert not missing, "برگه‌های نمونه بدون واترمارک: %s" % missing
