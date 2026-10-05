# -*- coding: utf-8 -*-
# T65 - تطابق فرمول‌های زندهٔ Excel میان خروجی Python و TypeScript.
#
# صادرکنندهٔ Python (xlsxwriter) سه فرمول =SUM زنده می‌نویسد:
#   1. جمع مبالغ فصول در برگهٔ «خلاصه مالی» (ستون C)
#   2. جمع کل بهای کل به‌هنگام در برگهٔ «صورت وضعیت» (ستون J)
#   3. جمع کل مقدار در برگهٔ «صورت وضعیت» (ستون F)
# این تست آن سلول‌ها را از خروجی واقعی استخراج می‌کند (فرمول + مقدار
# کش‌شده + قالب عدد) و در golden ثبت می‌کند تا نسخهٔ TS دقیقاً به آن برسد.
# مقدار کش‌شده (v) باید برابر با جمع واقعی سلول‌های بازه باشد.
import io
import os
import re
import sys

import openpyxl
import pytest

pytestmark = pytest.mark.slow

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from app import (DynamoMappingParser, ExcelExporter, RevitScheduleParser,
                 TechnicalOfficeEngine, _default_fehrest_db, default_settings)

from helper import golden_path, load_json, record  # noqa: E402

ROOT = os.path.join(os.path.dirname(__file__), "..", "..")
REVIT_DIR = os.path.join(ROOT, "data", "export revit")
DYNAMO_DIR = os.path.join(ROOT, "data", "dynamo")


def _real_result():
    tables = []
    for fname in sorted(os.listdir(REVIT_DIR)):
        if fname.endswith(".txt"):
            with open(os.path.join(REVIT_DIR, fname), "r", encoding="utf-8-sig") as fh:
                tables.append(RevitScheduleParser.parse(fname, fh.read()))
    results = []
    for root, _dirs, fnames in os.walk(DYNAMO_DIR):
        for fn in sorted(fnames):
            if fn.endswith(".txt"):
                with open(os.path.join(root, fn), "rb") as fh:
                    results.append(DynamoMappingParser.auto_detect_and_parse(fh.read(), fn))
    map_df = DynamoMappingParser.merge_formats(results)
    return TechnicalOfficeEngine(_default_fehrest_db(), default_settings()).compute(tables, map_df)


def _export_bytes(result):
    meta = {"project_name": "T65", "report_no": "1", "report_type": "موقت",
            "contract_no": "", "employer": "", "contractor": "", "consultant": "",
            "period_start": "", "period_end": "", "location": "", "contract_date": ""}
    exporter = ExcelExporter()
    return exporter.export(meta, result, default_settings(), [])


def _formula_cells(buf):
    """استخراج سلول‌های فرمول‌دار: {«برگه!آدرس»: {f, v, z}}."""
    wb = openpyxl.load_workbook(io.BytesIO(buf))                 # فرمول‌ها
    wbv = openpyxl.load_workbook(io.BytesIO(buf), data_only=True)  # مقادیر کش‌شده
    out = {}
    for ws in wb.worksheets:
        wsv = wbv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out[f"{ws.title}!{c.coordinate}"] = {
                        "f": c.value[1:],  # بدون «=» (مثل SheetJS)
                        "v": wsv[c.coordinate].value,
                        "z": c.number_format,
                    }
    return out


@pytest.fixture(scope="module")
def formulas():
    return _formula_cells(_export_bytes(_real_result()))


def test_T65_formula_cells_present(formulas):
    """هر سه فرمول زنده باید در خروجی واقعی وجود داشته باشند."""
    assert len(formulas) == 3, f"expected 3 formula cells, got {len(formulas)}"
    fs = sorted(formulas.values(), key=lambda x: x["f"])
    assert any(f["f"].startswith("SUM(C") for f in formulas.values()), "chapter SUM missing"
    assert any(f["f"].startswith("SUM(J") for f in formulas.values()), "boq total SUM missing"
    assert any(f["f"].startswith("SUM(F") for f in formulas.values()), "boq qty SUM missing"


def test_T65_cached_value_equals_range_sum(formulas):
    """مقدار کش‌شدهٔ هر فرمول باید برابر جمع سلول‌های بازه باشد."""
    buf = _export_bytes(_real_result())
    wbv = openpyxl.load_workbook(io.BytesIO(buf), data_only=True)
    for addr, info in formulas.items():
        sheet, coord = addr.rsplit("!", 1)
        m = re.match(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", info["f"])
        assert m, f"unparsed formula: {info['f']}"
        col1, r1, col2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        total = 0.0
        for r in range(r1, r2 + 1):
            v = wbv[sheet][f"{col1}{r}"].value
            if isinstance(v, (int, float)):
                total += v
        assert info["v"] == pytest.approx(total, abs=0.01), \
            f"{addr}: cached {info['v']} != range sum {total}"


def test_T65_golden(formulas):
    snap = {"formulas": {k: {kk: (round(vv, 6) if isinstance(vv, float) else vv)
                             for kk, vv in v.items()} for k, v in sorted(formulas.items())}}
    record("T65_excel_formulas", snap, generate=os.environ.get("GENERATE") == "1")
    golden = load_json(golden_path("T65_excel_formulas"))
    assert snap == golden, "formula cells changed — regenerate with GENERATE=1 if intended"


if __name__ == "__main__":
    fx = _formula_cells(_export_bytes(_real_result()))
    test_T65_formula_cells_present(fx)
    test_T65_cached_value_equals_range_sum(fx)
    print("T65 ok:", len(fx), "formula cells")
