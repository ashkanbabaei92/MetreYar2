# -*- coding: utf-8 -*-
# ─────────────────────────────────────────────────────────────
# T106 — فاز E-6 Part C گام ۵: ردیفِ «کار جدید» در خروجی Excel
#
# هر سه ناحیه‌ای که ردیفِ کار جدید باید ظاهر شود:
#   ۱) شیتِ «تعدیل پیمان» (جزئیاتِ فصول + تجهیز + کار جدید)
#   ۲) شیتِ «برآورد به‌هنگام» (جزئیاتِ فصول + تجهیز + کار جدید)
#   ۳) شیتِ «بخش قطعی و علی‌الحساب» (مبلغ + delta قطعی/علی‌الحساب)
#
# پاریتی با TypeScript: نسخهٔ TS (tests/char_ts/test_T106_e6_new_work_excel.ts)
# همان اعداد را از همان موتور می‌خواند. نکته: app.py مقادیر را با fa_float
# به‌صورتِ متنِ فارسیِ قالب‌بندی‌شده در سلول می‌نویسد، در حالی که TS سلولِ
# عددی می‌نویسد (تفاوتِ از‌پیش‌موجودِ هر دو پیاده‌سازی برای همهٔ ردیف‌ها)؛
# پاریتی روی اعداد بررسی می‌شود، نه روی بایت‌های سلول.
# ─────────────────────────────────────────────────────────────
import io
import os

import openpyxl
import pytest

from app import (DynamoMappingParser, ExcelExporter, RevitScheduleParser,
                 TechnicalOfficeEngine, _default_fehrest_db, default_settings,
                 fa_float)

ROOT = os.path.join(os.path.dirname(__file__), "..", "..")
REVIT_DIR = os.path.join(ROOT, "data", "export revit")
DYNAMO_DIR = os.path.join(ROOT, "data", "dynamo")

META = {"project_name": "T106", "report_no": "1", "report_type": "برآورد",
        "contract_no": "", "employer": "", "contractor": "", "consultant": "",
        "period_start": "", "period_end": "", "location": "",
        "contract_date": ""}

# اعداد مرجعِ کار جدید (رشتهٔ ابنیه، ۳-۲، مبلغ ۱e۹، مبنا ۱۴۰۴Q۲)
REF = {
    "amount": 883664623.308789,        # ۱e۹ × (۹۷۴۴٫۷/۱۱۶۰۸)/۰٫۹۵
    "io": 9744.7,
    "it_tadil": 12460.173698630137,
    "ratio_tadil": 1.2786616005244016,
    "value_tadil": 234449540.05814552,
    "it_baravord": 14333.1,
    "ratio_baravord": 1.470861083460753,
    "value_baravord": 1299747905.2559035,
}


def _settings():
    s = default_settings()
    s["tadil_enabled"] = True
    s["baravord_enabled"] = True
    s["proposal_date"] = "۱۴۰۳/۱۲/۲۹"
    s["work_period_start"] = "۱۴۰۴/۰۱/۰۱"
    s["work_period_end"] = "۱۴۰۴/۱۲/۲۹"
    s["new_item_method"] = "فهرست پایه"
    s["new_item_amount"] = 1_000_000_000.0
    s["new_item_base_period"] = "1404Q2"
    return s


def _real_result():
    tables = []
    for fname in sorted(os.listdir(REVIT_DIR)):
        if fname.endswith(".txt"):
            with open(os.path.join(REVIT_DIR, fname), "r", encoding="utf-8-sig") as fh:
                tables.append(RevitScheduleParser.parse(fname, fh.read()))
    results = []
    for root, _dirs, fnames in os.walk(DYNAMO_DIR):
        for fn in sorted(fnames):
            if fn.endswith(".txt"):
                with open(os.path.join(root, fn), "rb") as fh:
                    results.append(
                        DynamoMappingParser.auto_detect_and_parse(fh.read(), fn))
    return TechnicalOfficeEngine(
        _default_fehrest_db(), _settings()).compute(
            tables, DynamoMappingParser.merge_formats(results))


@pytest.fixture(scope="module")
def exported():
    res = _real_result()
    s = _settings()
    buf = ExcelExporter().export(META, res, s, [])
    wb = openpyxl.load_workbook(io.BytesIO(buf), data_only=True)
    rows = {}
    for sh in wb.sheetnames:
        for r in wb[sh].iter_rows():
            vals = [c.value for c in r]
            if vals and vals[0] == "کار جدید":
                rows[sh] = vals
    return rows, res["adjustment"]


# ── ۱) ردیفِ کار جدید در هر سه ناحیه هست ──

def test_T106_new_work_row_in_three_sheets(exported):
    rows, _ = exported
    assert set(rows) == {"تعدیل پیمان", "برآورد به‌هنگام", "بخش قطعی و علی‌الحساب"}


# ── ۲) شیتِ تعدیل پیمان: مقدار/شاخص‌ها/مبلغ تعدیل ──

def test_T106_tadil_detail_row(exported):
    rows, adj = exported
    nr = adj["tadil"]["new_work_row"]
    v = rows["تعدیل پیمان"]
    assert v[1] == "کار جدید (فهرست پایه) - بندهای ۳-۱..۳-۵"
    assert v[2] == fa_float(nr["amount"], 0)
    assert v[3] == fa_float(nr["io"], 1)
    assert v[4] == fa_float(nr["it"], 1)
    assert v[5] == fa_float(nr["ratio"], 4)
    assert v[6] == fa_float(nr["value"], 0)
    assert v[7] == nr["status"] + " - " + nr["src"]
    # اعدادِ مرجع
    assert nr["amount"] == pytest.approx(REF["amount"])
    assert nr["io"] == pytest.approx(REF["io"])
    assert nr["it"] == pytest.approx(REF["it_tadil"])
    assert nr["ratio"] == pytest.approx(REF["ratio_tadil"])
    assert nr["value"] == pytest.approx(REF["value_tadil"])


# ── ۳) شیتِ برآورد به‌هنگام: It برآوردی متفاوت ──

def test_T106_baravord_detail_row(exported):
    rows, adj = exported
    nr = adj["baravord"]["new_work_row"]
    v = rows["برآورد به‌هنگام"]
    assert v[1] == "کار جدید (فهرست پایه) - بندهای ۳-۱..۳-۵"
    assert v[2] == fa_float(nr["amount"], 0)
    assert v[3] == fa_float(nr["io"], 1)
    assert v[4] == fa_float(nr["it"], 1)
    assert v[5] == fa_float(nr["ratio"], 4)
    assert v[6] == fa_float(nr["value"], 0)
    assert nr["amount"] == pytest.approx(REF["amount"])
    assert nr["it"] == pytest.approx(REF["it_baravord"])
    assert nr["ratio"] == pytest.approx(REF["ratio_baravord"])
    assert nr["value"] == pytest.approx(REF["value_baravord"])


# ── ۴) شیتِ بخش قطعی و علی‌الحساب: تجمیع در td/tp ──

def test_T106_definitive_provisional_row(exported):
    rows, adj = exported
    nr = adj["tadil"]["new_work_row"]
    v = rows["بخش قطعی و علی‌الحساب"]
    assert v[2] == fa_float(nr["amount"], 0)
    assert v[3] == fa_float(nr["delta_definitive"], 0)
    assert v[4] == fa_float(nr["delta"] - nr["delta_definitive"], 0)
    assert v[5] == fa_float(nr["delta"], 0)
    assert v[6] == "همان منطقِ تجهیز (شاخص رشته)"
    assert nr["delta"] == pytest.approx(REF["value_tadil"])
    assert nr["delta_definitive"] == pytest.approx(REF["value_tadil"])


# ── ۵) بدون کار جدید، ردیفی در هیچ شیت نیست ──

def test_T106_no_new_work_no_row():
    tables = []
    for fname in sorted(os.listdir(REVIT_DIR)):
        if fname.endswith(".txt"):
            with open(os.path.join(REVIT_DIR, fname), "r", encoding="utf-8-sig") as fh:
                tables.append(RevitScheduleParser.parse(fname, fh.read()))
    results = []
    for root, _dirs, fnames in os.walk(DYNAMO_DIR):
        for fn in sorted(fnames):
            if fn.endswith(".txt"):
                with open(os.path.join(root, fn), "rb") as fh:
                    results.append(
                        DynamoMappingParser.auto_detect_and_parse(fh.read(), fn))
    s = default_settings()
    s["tadil_enabled"] = True
    s["baravord_enabled"] = True
    s["proposal_date"] = "۱۴۰۳/۱۲/۲۹"
    s["work_period_start"] = "۱۴۰۴/۰۱/۰۱"
    s["work_period_end"] = "۱۴۰۴/۱۲/۲۹"
    res = TechnicalOfficeEngine(
        _default_fehrest_db(), s).compute(
            tables, DynamoMappingParser.merge_formats(results))
    buf = ExcelExporter().export(META, res, s, [])
    wb = openpyxl.load_workbook(io.BytesIO(buf), data_only=True)
    hit = [sh for sh in wb.sheetnames
           for r in wb[sh].iter_rows()
           if r[0].value == "کار جدید"]
    assert hit == []
